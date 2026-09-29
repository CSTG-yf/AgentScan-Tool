from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "task" / "DefenseClaw_AgentRange靶场改造实施计划书_v1.0.md"
OUTPUT = ROOT / "task" / "DefenseClaw_AgentRange靶场改造实施计划书_v1.0.docx"

BLUE = "2E74B5"
NAVY = "1F4D78"
LIGHT = "F2F4F7"
PALE_BLUE = "EAF2F8"
GRAY = "5F6B76"
BLACK = "000000"
TABLE_WIDTH = 9360


def set_run_font(run, size=None, bold=None, color=None, east_asia="Microsoft YaHei"):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)


def widths_for(headers):
    n = len(headers)
    joined = "|".join(headers)
    if n == 2:
        return [3000, 6360]
    if n == 3 and "指标" in joined:
        return [2700, 1500, 5160]
    if n == 3:
        return [2200, 2300, 4860]
    if n == 4 and "风险" in joined:
        return [1600, 2200, 3100, 2460]
    if n == 4:
        return [1450, 1900, 2500, 3510]
    if n == 5:
        return [1150, 1600, 1200, 2650, 2760]
    base = TABLE_WIDTH // n
    result = [base] * n
    result[-1] += TABLE_WIDTH - sum(result)
    return result


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_inline_runs(paragraph, text, size=11, color=BLACK, bold=False):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=size, color=NAVY, east_asia="Microsoft YaHei")
            run.font.name = "Consolas"
            run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
            run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, bold=bold, color=color)


def add_page_field(paragraph):
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color=GRAY)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, NAVY, 8, 4),
    }
    for name, (size, color, before, after) in specs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_inline_runs(header, "DefenseClaw × AgentRange｜改造实施计划书 v1.0", size=9, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(footer)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("DEFENSECLAW 改造实施计划书")
    set_run_font(r, size=25, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("面向 AgentRange 靶场的需求、架构、任务分解与验收基线")
    set_run_font(r, size=14, color=BLUE)

    metadata = [
        ("版本", "v1.0"),
        ("日期", "2026-07-19"),
        ("范围", "DefenseClaw 主工程 + task/AgentRange-player-full"),
        ("状态", "实施基线"),
        ("验证约束", "不运行 Go 测试；不使用 CI/CD 或 GitHub 测试"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{label}：")
        set_run_font(r, size=10.5, bold=True, color=GRAY)
        r = p.add_run(value)
        set_run_font(r, size=10.5, color=BLACK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PALE_BLUE)
    p_pr.append(shd)
    add_inline_runs(
        p,
        "核心结论：DefenseClaw 仅增加可复用的身份上下文、运行时检查与决策能力；AgentRange 特有格式、证据、回放和指标全部由适配层承接。",
        size=11,
        color=NAVY,
        bold=True,
    )
    doc.add_page_break()


def add_table(doc, rows):
    headers = rows[0]
    table = doc.add_table(rows=len(rows), cols=len(headers))
    table.style = "Table Grid"
    widths = widths_for(headers)
    for row_idx, values in enumerate(rows):
        row = table.rows[row_idx]
        if row_idx == 0:
            set_repeat_header(row)
        for col_idx, value in enumerate(values):
            cell = row.cells[col_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            if row_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_cell_shading(cell, LIGHT)
                add_inline_runs(p, value, size=9.5, bold=True)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_inline_runs(p, value, size=9.2)
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    i = 0
    in_code = False
    code_lines = []
    started = False
    forced_breaks = {"## 6. 详细需求基线", "## 9. 实施任务分解与分配", "## 14. 交付物与验收证据"}
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        if not started:
            if line.startswith("## "):
                started = True
            else:
                i += 1
                continue

        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.22)
                p.paragraph_format.right_indent = Inches(0.12)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.line_spacing = 1.0
                p_pr = p._p.get_or_add_pPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "F7F8FA")
                p_pr.append(shd)
                for idx, code in enumerate(code_lines):
                    run = p.add_run(code + ("\n" if idx < len(code_lines) - 1 else ""))
                    set_run_font(run, size=8.5, color=BLACK, east_asia="Microsoft YaHei")
                    run.font.name = "Consolas"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
                    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|(?:\s*:?-+:?\s*\|)+$", lines[i + 1]):
            rows = []
            header = [c.strip() for c in line.strip("|").split("|")]
            rows.append(header)
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            add_table(doc, rows)
            continue

        if line.startswith("## "):
            if line in forced_breaks and len(doc.paragraphs) > 3:
                doc.add_page_break()
            p = doc.add_paragraph(style="Heading 1")
            add_inline_runs(p, line[3:], size=16, color=BLUE, bold=True)
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline_runs(p, line[4:], size=13, color=BLUE, bold=True)
        elif line.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            add_inline_runs(p, line[5:], size=12, color=NAVY, bold=True)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, line[2:])
        elif re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, re.sub(r"^\d+\.\s+", "", line))
        elif line.strip():
            p = doc.add_paragraph()
            add_inline_runs(p, line.strip())
        i += 1

    props = doc.core_properties
    props.title = "DefenseClaw 面向 AgentRange 靶场改造实施计划书 v1.0"
    props.subject = "需求基线、技术方案、任务分解、联调验证与交付验收"
    props.author = "DefenseClaw 项目组"
    props.keywords = "DefenseClaw, AgentRange, AI Agent Security, Implementation Plan"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
